from copy import deepcopy
from pathlib import Path

from docx import Document

SOURCE = Path("/Users/brittbowman/Documents/Client Documents/WoltersKluwer/Invoice_BB-2026-WK-002.docx")
OUTPUT = Path("/Users/brittbowman/Documents/DRIVE/helix-ai-demo/outputs/bloomreach-invoices")

INVOICES = [
    ("BB-2026-BR-001", "September 1, 2026", "September 16, 2026", "September 2026"),
    ("BB-2026-BR-002", "October 1, 2026", "October 16, 2026", "October 2026"),
    ("BB-2026-BR-003", "November 1, 2026", "November 16, 2026", "November 2026"),
    ("BB-2026-BR-004", "December 1, 2026", "December 16, 2026", "December 2026"),
]


def replace_paragraph(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def set_cell_paragraph(cell, index, text):
    replace_paragraph(cell.paragraphs[index], text)


def build(invoice_number, invoice_date, due_date, service_month):
    document = Document(SOURCE)

    header = document.tables[0]
    set_cell_paragraph(header.cell(0, 1), 1, f"Invoice #: {invoice_number}")
    set_cell_paragraph(header.cell(0, 1), 2, f"Invoice Date: {invoice_date}")
    set_cell_paragraph(header.cell(0, 1), 3, f"Due Date: {due_date}")

    billing = document.tables[1]
    set_cell_paragraph(billing.cell(0, 0), 1, "Bloomreach, Inc.")
    set_cell_paragraph(billing.cell(0, 0), 2, "Dept: Accounts Payable")
    set_cell_paragraph(billing.cell(0, 0), 3, "payables@bloomreach.com")
    set_cell_paragraph(billing.cell(0, 0), 4, "888-263-3917")
    set_cell_paragraph(billing.cell(0, 0), 5, "")
    set_cell_paragraph(billing.cell(0, 1), 0, "Supplier POC: Britt Bowman")
    set_cell_paragraph(billing.cell(0, 1), 1, "britt@brittbowman.ai")
    set_cell_paragraph(billing.cell(0, 1), 2, "Bloomreach POC: Amanda Cole | 713-882-0873")
    set_cell_paragraph(billing.cell(0, 1), 3, "amanda.cole@bloomreach.com")

    charges = document.tables[2]
    set_cell_paragraph(charges.cell(1, 0), 0, f"Fractional COO for Marketing — Monthly Retainer ({service_month})")
    set_cell_paragraph(charges.cell(1, 1), 0, "1")
    set_cell_paragraph(charges.cell(1, 2), 0, "$15,000.00")
    set_cell_paragraph(charges.cell(1, 3), 0, "$15,000.00")
    set_cell_paragraph(charges.cell(2, 3), 0, "$15,000.00 USD")

    target = OUTPUT / f"Invoice_{invoice_number}.docx"
    document.save(target)
    print(target)


OUTPUT.mkdir(parents=True, exist_ok=True)
for invoice in INVOICES:
    build(*invoice)
