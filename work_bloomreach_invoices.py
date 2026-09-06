from docx import Document

SOURCE = "/Users/brittbowman/Documents/Client Documents/WoltersKluwer/Invoice_BB-2026-WK-002.docx"

document = Document(SOURCE)
print("PARAGRAPHS")
for index, paragraph in enumerate(document.paragraphs):
    print(index, repr(paragraph.text), [repr(run.text) for run in paragraph.runs])

print("TABLES")
for table_index, table in enumerate(document.tables):
    for row_index, row in enumerate(table.rows):
        for cell_index, cell in enumerate(row.cells):
            print("TABLE", table_index, "ROW", row_index, "CELL", cell_index, repr(cell.text))
            for paragraph_index, paragraph in enumerate(cell.paragraphs):
                print("  P", paragraph_index, repr(paragraph.text), [repr(run.text) for run in paragraph.runs])
